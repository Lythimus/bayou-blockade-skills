#!/usr/bin/env python3
"""Chemical incident records: U.S. Coast Guard National Response Center (NRC) spill and release reports,
U.S. Chemical Safety Board (CSB) investigations, and CSB Accidental Release Events (40 CFR Part 1604).

Sources:
  NRC  yearly FOIA workbooks, https://nrc.uscg.mil/FOIAFiles/CY{YY}.xlsx (1990 to the current year,
       10-18 MB each; the current year is re-published as reports come in)
  CSB  investigation pages, https://www.csb.gov/investigations/
  CSB  Accidental Release Events workbook, linked from https://www.csb.gov/news/incident-report-rule-form-/

Run with: uv run -q --no-project --with openpyxl python -I hazmat.py <command> ...
Cache: ~/.cache/bayou/csb-nrc-hazmat/ (NRC years parsed to compact JSON; the current and previous year are
re-checked against Last-Modified daily, older years kept; CSB pages 7 days).

  hazmat.py health
  hazmat.py nrc --state LA [--county X] [--city X] [--company "A|B"] [--text X] [--material X]
                [--years 2021-2026] [--near LAT LON --radius-km N] [--csv out.csv] [--yes]
  hazmat.py nrc-report <SEQNOS> [--year YYYY]
  hazmat.py nrc-summary --state LA [same filters as nrc]
  hazmat.py csb-investigations [--state LA] [--company X] [--status current|completed]
  hazmat.py csb-detail <slug or URL>
  hazmat.py csb-releases [--state Louisiana] [--company X] [--city X] [--csv out.csv]

Exit codes: 3 not found, 4 download needs confirmation (rerun with --yes), 5 source down.
"""
import argparse
import collections
import csv
import datetime
import email.utils
import gzip
import html
import io
import json
import math
import os
import re
import sys
import time
import urllib.error
import urllib.parse
import urllib.request

CACHE = os.path.expanduser("~/.cache/bayou/csb-nrc-hazmat")
TTL = 7 * 24 * 3600
UA = "Mozilla/5.0 (bayou-csb-nrc-hazmat)"
NRC = "https://nrc.uscg.mil/"
NRC_FILE = NRC + "FOIAFiles/CY%02d.xlsx"
NRC_FIRST = 1990
CSB = "https://www.csb.gov"
CSB_LISTS = {"completed": CSB + "/investigations/completed-investigations/?Type=2&pg=%d",
             "current": CSB + "/investigations/current-investigations/?Type=1&pg=%d"}
ARRE_PAGE = CSB + "/news/incident-report-rule-form-/"
# A guarded download is anything over one uncached workbook or this many MB.
GUARD_MB = 50
NRC_NOTE = ("NRC is a call center, not a response agency; in its words the posted spreadsheets 'contain INITIAL "
            "incident data that has not been validated or investigated by a federal/state response agency'. "
            "'UNKNOWN AMOUNT' (amount 0) means the caller gave no quantity, not zero; "
            "CAS 000000-00-0 means none was given. A report is not proof a reportable quantity was exceeded, and "
            "one event can produce several reports (updates, PHMSA 48-hour follow-ups, multiple callers).")
ARRE_NOTE = ("40 CFR 1604.3 requires a stationary source's owner or operator to report an accidental release to the "
             "ambient air that results in a fatality, serious injury (death or inpatient hospitalization) or "
             "substantial property damage (>= $1,000,000), effective 2020-03-23. Either the NRC report number goes to "
             "CSB within 30 minutes of the NRC call (so an event may also have an NRC report), or a report goes "
             "directly to CSB within 8 hours. CSB revises the list quarterly, including late-reported events. A "
             "listing does not mean CSB investigated; full reports are available by FOIA (40 CFR 1604.6).")
STATES = {"AL": "Alabama", "AK": "Alaska", "AZ": "Arizona", "AR": "Arkansas", "CA": "California", "CO": "Colorado",
          "CT": "Connecticut", "DE": "Delaware", "DC": "District of Columbia", "FL": "Florida", "GA": "Georgia",
          "HI": "Hawaii", "ID": "Idaho", "IL": "Illinois", "IN": "Indiana", "IA": "Iowa", "KS": "Kansas",
          "KY": "Kentucky", "LA": "Louisiana", "ME": "Maine", "MD": "Maryland", "MA": "Massachusetts",
          "MI": "Michigan", "MN": "Minnesota", "MS": "Mississippi", "MO": "Missouri", "MT": "Montana",
          "NE": "Nebraska", "NV": "Nevada", "NH": "New Hampshire", "NJ": "New Jersey", "NM": "New Mexico",
          "NY": "New York", "NC": "North Carolina", "ND": "North Dakota", "OH": "Ohio", "OK": "Oklahoma",
          "OR": "Oregon", "PA": "Pennsylvania", "RI": "Rhode Island", "SC": "South Carolina", "SD": "South Dakota",
          "TN": "Tennessee", "TX": "Texas", "UT": "Utah", "VT": "Vermont", "VA": "Virginia", "WA": "Washington",
          "WV": "West Virginia", "WI": "Wisconsin", "WY": "Wyoming", "PR": "Puerto Rico", "GU": "Guam",
          "VI": "Virgin Islands", "AS": "American Samoa", "MP": "Northern Mariana Islands"}


class SourceDown(Exception):
    pass


class NeedsConfirm(Exception):
    pass


def die(msg, code=2):
    print(msg, file=sys.stderr)
    sys.exit(code)


def today():
    return datetime.date.today().isoformat()


def note(msg):
    print(msg, file=sys.stderr)


# ---------------------------------------------------------------- HTTP and cache

def _req(url, method="GET", timeout=120):
    """Response object; retries the NRC server's intermittent 503s."""
    req = urllib.request.Request(url, method=method, headers={"User-Agent": UA})
    for attempt in range(4):
        try:
            return urllib.request.urlopen(req, timeout=timeout)
        except urllib.error.HTTPError as e:
            if e.code == 404:
                return None
            if e.code in (429, 502, 503, 504) and attempt < 3:
                time.sleep(2 + 3 * attempt)
                continue
            raise SourceDown("%s: HTTP %s" % (url, e.code))
        except (urllib.error.URLError, TimeoutError, ConnectionError) as e:
            if attempt < 3:
                time.sleep(2)
                continue
            raise SourceDown("%s: %s" % (url, e))


def get(url, timeout=120):
    r = _req(url, timeout=timeout)
    if r is None:
        return None
    with r:
        return r.read()


def head(url):
    """(size bytes, Last-Modified as ISO date, raw Last-Modified) or None when the file does not exist."""
    r = _req(url, method="HEAD", timeout=60)
    if r is None:
        return None
    with r:
        lm = r.headers.get("Last-Modified")
        size = int(r.headers.get("Content-Length") or 0)
    iso = email.utils.parsedate_to_datetime(lm).date().isoformat() if lm else None
    return size, iso, lm


def cached_text(name, url, ttl=TTL, refresh=False):
    path = os.path.join(CACHE, name)
    if os.path.exists(path) and not refresh and time.time() - os.path.getmtime(path) < ttl:
        with open(path, encoding="utf8") as fh:
            return fh.read()
    try:
        body = get(url)
    except SourceDown:
        if os.path.exists(path):
            note("note: %s unreachable; using cached copy" % url)
            with open(path, encoding="utf8") as fh:
                return fh.read()
        raise
    if body is None:
        return None
    text = body.decode("utf8", "replace")
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path + ".part", "w", encoding="utf8") as fh:
        fh.write(text)
    os.replace(path + ".part", path)
    return text


def norm(s):
    return re.sub(r"[^a-z0-9 ]", " ", (s or "").lower()).split()


def words_in(needle, hay):
    hw = set(norm(hay))
    return all(w in hw for w in norm(needle))


def haversine_km(a, b, c, d):
    p1, p2 = math.radians(a), math.radians(c)
    dp, dl = p2 - p1, math.radians(d - b)
    h = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * 6371.0 * math.asin(math.sqrt(h))


# ---------------------------------------------------------------- NRC workbooks

def nrc_dir():
    d = os.path.join(CACHE, "nrc")
    os.makedirs(d, exist_ok=True)
    return d


def nrc_index():
    p = os.path.join(nrc_dir(), "index.json")
    if os.path.exists(p):
        with open(p) as fh:
            return json.load(fh)
    return {}


def save_index(idx):
    p = os.path.join(nrc_dir(), "index.json")
    with open(p + ".part", "w") as fh:
        json.dump(idx, fh, indent=1, sort_keys=True)
    os.replace(p + ".part", p)


def parse_years(spec):
    cur = datetime.date.today().year
    if not spec:
        return list(range(cur - 4, cur + 1))
    out = []
    for part in spec.split(","):
        if "-" in part:
            a, b = part.split("-")
            out += range(int(a), int(b) + 1)
        else:
            out.append(int(part))
    bad = [y for y in out if y < NRC_FIRST or y > cur]
    if bad:
        die("NRC yearly files cover %d-%d; got %s" % (NRC_FIRST, cur, bad))
    return sorted(set(out))


def stale_check_needed(year, entry, refresh):
    """Only the current and previous year change routinely. NRC re-posts older years rarely (CY90, CY10, CY22
    and CY23 all carry Last-Modified 2024-04-18), so those are re-checked only on --refresh."""
    if entry is None:
        return True
    if year < datetime.date.today().year - 1:
        return refresh
    return refresh or time.time() - entry.get("checked", 0) > 24 * 3600


def plan_years(years, refresh):
    """[(year, url, size, last_modified_iso, raw_lm)] needing download, with HEAD done."""
    idx = nrc_index()
    todo = []
    for y in years:
        entry = idx.get(str(y))
        if entry and not os.path.exists(os.path.join(nrc_dir(), "CY%d.json.gz" % y)):
            entry = None
        if not stale_check_needed(y, entry, refresh):
            continue
        url = NRC_FILE % (y % 100)
        try:
            h = head(url)
        except SourceDown:
            if entry:
                note("note: %s unreachable; using cached CY%d (published %s)" % (url, y, entry.get("last_modified")))
                continue
            raise
        if h is None:
            if entry:
                continue
            die("NRC has no file for %d (%s)" % (y, url), 3)
        size, iso, raw = h
        if entry and entry.get("last_modified_raw") == raw:
            entry["checked"] = time.time()
            idx[str(y)] = entry
            continue
        todo.append((y, url, size, iso, raw))
    save_index(idx)
    return todo


def guard(todo, yes):
    if not todo:
        return
    mb = sum(t[2] for t in todo) / 1e6
    lines = ["  CY%d  %.1f MB  (published %s)  %s" % (t[0], t[2] / 1e6, t[3], t[1]) for t in todo]
    msg = "NRC download: %d workbook(s), %.1f MB total\n%s" % (len(todo), mb, "\n".join(lines))
    if not yes and (len(todo) > 1 or mb > GUARD_MB):
        raise NeedsConfirm(msg + "\nOver the 1-file / %d MB limit. Ask the user, then rerun with --yes, "
                                 "or narrow --years." % GUARD_MB)
    note(msg)


def dms(deg, mn, sec, quad):
    try:
        v = float(deg) + float(mn or 0) / 60 + float(sec or 0) / 3600
    except (TypeError, ValueError):
        return None
    if v == 0:
        return None
    return -v if (quad or "").upper() in ("S", "W") else v


def clean(v):
    if v is None:
        return None
    if isinstance(v, datetime.datetime):
        return v.isoformat(sep=" ", timespec="minutes")
    if isinstance(v, (datetime.date, datetime.time)):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return int(v)
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def parse_workbook(path):
    """{SEQNOS: {sheet: [row dict, ...]}} with empty cells dropped; sheets differ in row counts per report."""
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    recs = collections.defaultdict(lambda: collections.defaultdict(list))
    for ws in wb:
        it = ws.iter_rows(values_only=True)
        try:
            hdr = next(it)
        except StopIteration:
            continue
        if not hdr or "SEQNOS" not in hdr:
            continue
        hdr = [str(h).strip() if h else None for h in hdr]
        k = hdr.index("SEQNOS")
        last = None
        for r in it:
            seq = r[k] if k < len(r) else None
            if seq in (None, ""):
                continue
            try:
                seq = int(seq)
            except (TypeError, ValueError):
                # A line break inside a free-text field spills the rest of that record onto a new row whose
                # first cell is text, not a SEQNOS. Keep the fragment with the record it came from.
                if last is not None:
                    frag = " | ".join(str(clean(v)) for v in r if clean(v) not in (None, ""))
                    last.setdefault("_CONTINUED", []).append(frag)
                continue
            row = {h: clean(v) for h, v in zip(hdr, r) if h and h != "SEQNOS" and clean(v) not in (None, "")}
            recs[seq][ws.title].append(row)
            last = row
    wb.close()
    return {s: dict(v) for s, v in recs.items()}


def download_year(y, url, size, iso, raw):
    note("downloading CY%d (%.1f MB) ..." % (y, size / 1e6))
    body = get(url, timeout=600)
    if body is None:
        die("NRC file disappeared: %s" % url, 3)
    if size and len(body) != size:
        raise SourceDown("%s: got %d of %d bytes" % (url, len(body), size))
    # openpyxl refuses files that do not end in .xlsx.
    tmp = os.path.join(nrc_dir(), "CY%d.tmp.xlsx" % y)
    with open(tmp, "wb") as fh:
        fh.write(body)
    try:
        recs = parse_workbook(tmp)
    finally:
        os.remove(tmp)
    out = os.path.join(nrc_dir(), "CY%d.json.gz" % y)
    with gzip.open(out + ".part", "wt", encoding="utf8") as fh:
        json.dump(recs, fh, separators=(",", ":"), default=str)
    os.replace(out + ".part", out)
    idx = nrc_index()
    seqs = [int(s) for s in recs]
    idx[str(y)] = {"url": url, "size": size, "last_modified": iso, "last_modified_raw": raw,
                   "retrieved": today(), "checked": time.time(), "reports": len(recs),
                   "min_seqnos": min(seqs) if seqs else None, "max_seqnos": max(seqs) if seqs else None}
    save_index(idx)


def ensure_years(years, refresh, yes):
    todo = plan_years(years, refresh)
    guard(todo, yes)
    for t in todo:
        download_year(*t)


def load_year(y):
    with gzip.open(os.path.join(nrc_dir(), "CY%d.json.gz" % y), "rt", encoding="utf8") as fh:
        return json.load(fh)


def first(rec, sheet):
    rows = rec.get(sheet) or []
    return rows[0] if rows else {}


# Follow-up reports cite the original: "UPDATE TO NRC 1461720", "CORRECTION TO NRC REPORT 1470627",
# "48HR PHMSA UPDATE TO INCIDENT REPORT # 1453392".
def iso_dt(v):
    """Workbooks before ~2020 hold real datetimes; later ones hold 'M/D/YYYY H:MM' strings."""
    if not isinstance(v, str):
        return v
    m = re.match(r"(\d{1,2})/(\d{1,2})/(\d{4})(?:\s+(\d{1,2}):(\d{2}))?", v)
    if not m:
        return v
    mo, d, y, hh, mm = m.groups()
    return "%s-%02d-%02d%s" % (y, int(mo), int(d), " %02d:%s" % (int(hh), mm) if hh else "")


UPDATE_RE = re.compile(r"(?:UPDATE|CORRECTION|FOLLOW[- ]?UP)S?\s+(?:TO|FOR|OF)\s+(?:THE\s+)?(?:NRC\s*)?"
                       r"(?:INCIDENT\s*)?(?:REPORT\s*)?(?:NO\.?|NUMBER|#)?\s*(\d{6,8})", re.I)


def summarize(seq, rec, year):
    c, ic, d = first(rec, "CALLS"), first(rec, "INCIDENT_COMMONS"), first(rec, "INCIDENT_DETAILS")
    mats = []
    for m in rec.get("MATERIAL_INVOLVED", []):
        unit = m.get("UNIT_OF_MEASURE")
        amt = m.get("AMOUNT_OF_MATERIAL")
        mats.append({"name": m.get("NAME_OF_MATERIAL"),
                     "cas": None if m.get("CAS_NUMBER") in (None, "000000-00-0") else m.get("CAS_NUMBER"),
                     "amount": None if unit == "UNKNOWN AMOUNT" or not amt else amt,
                     "unit": None if unit == "UNKNOWN AMOUNT" else unit,
                     "reached_water": m.get("IF_REACHED_WATER"),
                     "amount_in_water": m.get("AMOUNT_IN_WATER") or None,
                     "unit_in_water": m.get("UNIT_OF_MEASURE_REACH_WATER")})
    for m in rec.get("MATERIAL_INV0LVED_CR", []):
        mats.append({"name": m.get("NAME_OF_MATERIAL"), "cas": m.get("CAS_NUMBER"),
                     "continuous_upper_bound": m.get("UPPER_BOUNDS"), "unit": m.get("UPPER_BOUNDS_UNIT"),
                     "rate": m.get("UPPER_BOUNDS_RATE")})
    lat = dms(ic.get("LAT_DEG"), ic.get("LAT_MIN"), ic.get("LAT_SEC"), ic.get("LAT_QUAD"))
    lon = dms(ic.get("LONG_DEG"), ic.get("LONG_MIN"), ic.get("LONG_SEC"), ic.get("LONG_QUAD"))
    desc = ic.get("DESCRIPTION_OF_INCIDENT") or ""
    upd = UPDATE_RE.search(desc)
    return {
        "seqnos": int(seq), "year": year,
        "received": iso_dt(c.get("DATE_TIME_RECEIVED")), "incident_date": iso_dt(ic.get("INCIDENT_DATE_TIME")),
        "date_kind": ic.get("INCIDENT_DTG"), "calltype": c.get("CALLTYPE"),
        "type": ic.get("TYPE_OF_INCIDENT"), "cause": ic.get("INCIDENT_CAUSE"),
        "potential": ic.get("POTENTIAL_FLAG") == "Y",
        "company": c.get("RESPONSIBLE_COMPANY"), "org_type": c.get("RESPONSIBLE_ORG_TYPE"),
        "location": ic.get("INCIDENT_LOCATION"),
        "address": " ".join(x for x in (ic.get("LOCATION_ADDRESS"), ic.get("LOCATION_STREET1"),
                                        ic.get("LOCATION_STREET2")) if x) or None,
        "city": ic.get("LOCATION_NEAREST_CITY"), "county": ic.get("LOCATION_COUNTY"),
        "state": ic.get("LOCATION_STATE"), "zip": ic.get("LOCATION_ZIP"),
        "lat": round(lat, 5) if lat else None, "lon": round(lon, 5) if lon else None,
        "description": desc,
        "update_of": int(upd.group(1)) if upd else None,
        "materials": mats,
        "medium": d.get("MEDIUM_DESC"), "body_of_water": d.get("BODY_OF_WATER"),
        "injured": d.get("NUMBER_INJURED"), "hospitalized": d.get("NUMBER_HOSPITALIZED"),
        "fatalities": d.get("NUMBER_FATALITIES"), "evacuated": d.get("NUMBER_EVACUATED"),
        "fire": d.get("FIRE_INVOLVED"), "release_secured": d.get("RELEASE_SECURED"),
        "duration": d.get("ESTIMATED_DURATION_OF_RELEASE"), "rate": d.get("RELEASE_RATE"),
        "state_agency_on_scene": d.get("STATE_AGENCY_ON_SCENE"),
        "state_agency_report": d.get("STATE_AGENCY_REPORT_NUM"),
    }


def norm_county(s):
    s = (s or "").upper().replace(".", " ")
    s = re.sub(r"\b(PARISH|COUNTY|BOROUGH)\b", " ", s)
    s = re.sub(r"\bSAINT\b", "ST", s)
    for full, ab in (("EAST", "E"), ("WEST", "W"), ("NORTH", "N"), ("SOUTH", "S")):
        s = re.sub(r"^%s\b" % full, ab, s.strip())
    return " ".join(s.split())


def nrc_filter(a, years):
    st = a.state.upper()
    county = norm_county(a.county) if a.county else None
    companies = [x for x in (a.company or "").split("|") if x.strip()]
    rows, no_coords = [], 0
    for y in years:
        for seq, rec in load_year(y).items():
            ic = first(rec, "INCIDENT_COMMONS")
            if (ic.get("LOCATION_STATE") or "").upper() != st:
                continue
            calltype = first(rec, "CALLS").get("CALLTYPE")
            if calltype != "INC" and not a.include_drills:
                continue
            if county and norm_county(ic.get("LOCATION_COUNTY")) != county:
                continue
            if a.city and not words_in(a.city, ic.get("LOCATION_NEAREST_CITY")):
                continue
            if a.type and not words_in(a.type, ic.get("TYPE_OF_INCIDENT")):
                continue
            s = summarize(seq, rec, y)
            if companies and not any(words_in(x, s["company"]) for x in companies):
                continue
            if a.text and not words_in(a.text, " ".join(filter(None, (s["description"], s["location"], s["address"])))):
                continue
            if a.material and not any(words_in(a.material, m.get("name")) or a.material == m.get("cas")
                                      for m in s["materials"]):
                continue
            if a.near:
                if s["lat"] is None:
                    no_coords += 1
                    continue
                dist = haversine_km(a.near[0], a.near[1], s["lat"], s["lon"])
                if dist > a.radius_km:
                    continue
                s["distance_km"] = round(dist, 2)
            rows.append(s)
    rows.sort(key=lambda r: r["seqnos"], reverse=True)
    return rows, no_coords


def nrc_vintage(years):
    idx = nrc_index()
    return [{"year": y, "published": idx.get(str(y), {}).get("last_modified"),
             "retrieved": idx.get(str(y), {}).get("retrieved"), "url": NRC_FILE % (y % 100)} for y in years]


def filters_desc(a):
    parts = ["state=%s" % a.state]
    for k in ("county", "city", "type", "company", "text", "material"):
        if getattr(a, k):
            parts.append("%s=%s" % (k, getattr(a, k)))
    if a.near:
        parts.append("within %g km of %g,%g" % (a.radius_km, a.near[0], a.near[1]))
    return ", ".join(parts)


def cmd_nrc(a):
    if a.near and not a.radius_km:
        die("--near needs --radius-km")
    years = parse_years(a.years)
    ensure_years(years, a.refresh, a.yes)
    rows, no_coords = nrc_filter(a, years)
    if a.csv:
        cols = ["seqnos", "year", "received", "incident_date", "type", "cause", "potential", "company",
                "address", "city", "county", "state", "lat", "lon", "medium", "body_of_water", "injured",
                "hospitalized", "fatalities", "evacuated", "update_of", "materials", "description"]
        with open(a.csv, "w", newline="", encoding="utf8") as fh:
            w = csv.writer(fh)
            w.writerow(cols)
            for r in rows:
                mats = "; ".join("%s %s %s" % (m.get("name") or "?", m.get("amount") or "unknown amount",
                                               m.get("unit") or "") for m in r["materials"])
                w.writerow([mats if c == "materials" else r.get(c) for c in cols])
    if not rows:
        die("no NRC reports match: %s (years %s%s)" % (filters_desc(a), ", ".join(map(str, years)),
            "; %d matches lacked coordinates" % no_coords if no_coords else ""), 3)
    return {"filters": filters_desc(a), "count": len(rows), "rows": rows, "no_coords_skipped": no_coords,
            "updates": sum(1 for r in rows if r["update_of"]), "potential": sum(1 for r in rows if r["potential"]),
            "csv": a.csv, "vintage": nrc_vintage(years), "note": NRC_NOTE}


def cmd_nrc_summary(a):
    a.csv = None
    years = parse_years(a.years)
    ensure_years(years, a.refresh, a.yes)
    rows, no_coords = nrc_filter(a, years)
    if not rows:
        die("no NRC reports match: %s" % filters_desc(a), 3)
    by = {k: collections.Counter() for k in ("year", "type", "medium", "company", "material", "county", "cause")}
    tot = collections.Counter()
    for r in rows:
        for k in ("year", "type", "medium", "company", "county", "cause"):
            by[k][r.get(k) or "(blank)"] += 1
        for m in {(m.get("name") or "(blank)") for m in r["materials"]}:
            by["material"][m] += 1
        for k in ("injured", "hospitalized", "fatalities", "evacuated"):
            try:
                tot[k] += int(r.get(k) or 0)
            except (TypeError, ValueError):
                pass
        tot["with_injury"] += bool(r.get("injured"))
        tot["with_evacuation"] += bool(r.get("evacuated"))
        tot["reached_water"] += any(m.get("reached_water") == "YES" for m in r["materials"])
    return {"filters": filters_desc(a), "count": len(rows), "updates": sum(1 for r in rows if r["update_of"]),
            "potential": sum(1 for r in rows if r["potential"]), "no_coords_skipped": no_coords,
            # Company keeps a long tail: one-off spellings are how aliases get found.
            "by": {k: v.most_common({"company": 60, "material": 25, "cause": 15}.get(k)) for k, v in by.items()},
            "totals": dict(tot), "vintage": nrc_vintage(years), "note": NRC_NOTE}


def year_for_seqnos(seq):
    idx = nrc_index()
    for y, e in idx.items():
        if e.get("min_seqnos") and e["min_seqnos"] <= seq <= e["max_seqnos"]:
            return int(y)
    # SEQNOS run in call order, so the nearest cached years bracket the right one.
    below = [(e["max_seqnos"], int(y)) for y, e in idx.items() if e.get("max_seqnos") and e["max_seqnos"] < seq]
    above = [(e["min_seqnos"], int(y)) for y, e in idx.items() if e.get("min_seqnos") and e["min_seqnos"] > seq]
    if below and above and max(below)[1] + 1 == min(above)[1] - 1:
        return max(below)[1] + 1
    return None


def cmd_nrc_report(a):
    seq = int(a.seqnos)
    y = a.year or year_for_seqnos(seq)
    if y is None:
        die("cannot tell which year SEQNOS %d falls in from the cached years (%s). Pass --year YYYY; the report "
            "is filed under the year NRC received the call." % (seq, ", ".join(sorted(nrc_index())) or "none"), 2)
    ensure_years([y], a.refresh, a.yes)
    rec = load_year(y).get(str(seq))
    if rec is None:
        die("SEQNOS %d is not in the CY%d workbook" % (seq, y), 3)
    return {"seqnos": seq, "year": y, "summary": summarize(seq, rec, y), "sheets": rec,
            "vintage": nrc_vintage([y]), "note": NRC_NOTE}


# ---------------------------------------------------------------- CSB investigations

def para(s):
    """Tags stripped, runs of blank lines and indentation collapsed."""
    return re.sub(r"\s*\n\s*", "\n", strip_tags(s)).strip()


def strip_tags(s):
    s = re.sub(r"<br\s*/?>|</p>", "\n", s or "", flags=re.I)
    return re.sub(r"[ \t]+", " ", html.unescape(re.sub(r"<[^>]+>", "", s))).strip()


CARD_RE = re.compile(r'<a href="(/[^"]+/)"[^>]*class="linkHd"[^>]*><b>(.*?)</b></a>(.*?)class="rmore"', re.S)


def csb_cards(status, refresh):
    out, seen = [], set()
    for pg in range(1, 60):
        text = cached_text("csb/%s_%d.html" % (status, pg), CSB_LISTS[status] % pg, refresh=refresh)
        if text is None:
            break
        new = 0
        for m in CARD_RE.finditer(text):
            slug, title, body = m.group(1), strip_tags(m.group(2)), m.group(3)
            if slug in seen:
                continue
            seen.add(slug)
            new += 1
            labels = dict((k.rstrip(":").strip(), strip_tags(v)) for k, v in
                          re.findall(r'_lbl\w+"><b>([^<]*)</b>(.*?)</span>', body, re.S))
            desc = re.search(r'_lblDesc">(.*?)</span>', body, re.S)
            date_label = next((k for k in labels if k != "Location"), None)
            out.append({"title": title, "slug": slug.strip("/"), "url": CSB + slug, "status": status,
                        "location": re.sub(r"\s+,", ",", labels.get("Location", "")),
                        "date_label": date_label, "date": labels.get(date_label) if date_label else None,
                        "summary": " ".join(strip_tags(desc.group(1)).split()) + "..." if desc else None})
        if not new:
            break
    return out


def cmd_csb_investigations(a):
    statuses = [a.status] if a.status else ["current", "completed"]
    rows = []
    for st in statuses:
        rows += csb_cards(st, a.refresh)
    if a.state:
        code = a.state.upper() if len(a.state) == 2 else next((k for k, v in STATES.items()
                                                               if v.lower() == a.state.lower()), a.state)
        rows = [r for r in rows if re.search(r",\s*%s\b" % re.escape(code), r["location"], re.I)]
    if a.company:
        rows = [r for r in rows if words_in(a.company, r["title"] + " " + (r["summary"] or ""))]
    if not rows:
        die("no CSB investigations match", 3)
    return {"count": len(rows), "rows": rows, "source": CSB + "/investigations/", "retrieved": today()}


DETAIL_RE = re.compile(r"<strong class=\"bold\">([^<]+?):\s*</strong>(.*?)(?=<span class=\"sep\">|</p>)", re.S)


def cmd_csb_detail(a):
    slug = urllib.parse.urlparse(a.slug).path.strip("/") if a.slug.startswith("http") else a.slug.strip("/")
    url = "%s/%s/" % (CSB, slug)
    text = cached_text("csb/detail/%s.html" % slug, url, refresh=a.refresh)
    if text is None or "Accident Description" not in text:
        # CSB slugs often end in a stray hyphen ("...-and-fire-/"), so offer the listed ones.
        key = set(norm(slug.replace("-", " ")))
        near = [c["slug"] for c in csb_cards("completed", False) + csb_cards("current", False)
                if key and key <= set(norm(c["slug"].replace("-", " ")))]
        die("no CSB investigation page at %s%s" % (url, "; listed slugs: " + ", ".join(near) if near else
            " (run csb-investigations to find the slug)"), 3)
    title = strip_tags((re.search(r'<h1 class="title">(.*?)</h1>', text, re.S) or [None, ""])[1])
    facts = {}
    for k, v in DETAIL_RE.findall(text):
        v = strip_tags(re.sub(r'<span id="CT_Main_0_LblLocationLabel">.*?</span>', "", v))
        facts[k.strip()] = v.lstrip("</b>").strip()
    body = re.search(r'<div class="details">.*?</div>\s*<p>(.*?)</div>', text, re.S)
    # The main column repeats the final-report links inside an HTML comment; skip it.
    text_nc = re.sub(r"<!--.*?-->", "", text, flags=re.S)
    blocks = {}
    for m in re.finditer(r'<h2 class="title title2">([^<]+)</h2>(.*?)(?=<h2 class="title title2">|<footer|$)',
                         text_nc, re.S):
        name, chunk = m.group(1).strip(), m.group(2)
        if name == "Recommendations":
            continue
        links = []
        for href, label in re.findall(r'<a [^>]*href="([^"#][^"]*)"[^>]*>(.*?)</a>', chunk, re.S):
            label = strip_tags(label)
            if label and not href.startswith("javascript"):
                links.append({"label": label, "url": urllib.parse.urljoin(CSB, href)})
        if links:
            blocks[name] = links[:40]
    recs = []
    for li in re.findall(r'<li class="ac">(.*?)</li>', text_nc, re.S):
        to = strip_tags((re.search(r'class="head">(.*?)</a>', li, re.S) or [None, ""])[1])
        for sub in re.split(r'<div class="recSubItem">', li)[1:]:
            num = re.search(r'lblRecNumber">(.*?)</span>', sub, re.S)
            desc = re.search(r'lblDesc">(.*?)</span>\s*</div>', sub, re.S)
            stat = re.search(r'lnkStatus"[^>]*>(.*?)</a>', sub, re.S)
            recs.append({"recipient": to, "number": strip_tags(num.group(1)) if num else None,
                         "text": para(desc.group(1)) if desc else None,
                         "status": strip_tags(stat.group(1)) if stat else None})
    desc = para(body.group(1)) if body else None
    return {"title": title, "url": url, "facts": facts, "description": desc,
            "documents": blocks, "recommendations": recs, "retrieved": today()}


# ---------------------------------------------------------------- CSB Accidental Release Events (Part 1604)

def arre_link(refresh):
    text = cached_text("csb/arre_page.html", ARRE_PAGE, ttl=24 * 3600, refresh=refresh)
    if text is None:
        raise SourceDown("%s: 404" % ARRE_PAGE)
    m = re.search(r'href="([^"]*Accidental_Release_Events[^"]*\.xlsx)"', text, re.I)
    if not m:
        raise SourceDown("no Accidental Release Events workbook linked from %s" % ARRE_PAGE)
    vols = [urllib.parse.urljoin(CSB, h) for h in re.findall(r'href="([^"]*Incident_Reports?_Volume[^"]*\.pdf)"',
                                                             text, re.I)]
    return urllib.parse.urljoin(CSB, m.group(1)), vols


def arre_rows(refresh):
    url, vols = arre_link(refresh)
    name = os.path.basename(urllib.parse.urlparse(url).path).lower()
    path = os.path.join(CACHE, "csb", name)
    if refresh or not os.path.exists(path):
        body = get(url)  # urllib follows the 301 to CSB's lowercase asset path
        if body is None or not body.startswith(b"PK"):
            raise SourceDown("%s did not return a workbook" % url)
        os.makedirs(os.path.dirname(path), exist_ok=True)
        with open(path + ".part", "wb") as fh:
            fh.write(body)
        os.replace(path + ".part", path)
    import openpyxl
    # Report hyperlinks are only exposed when the workbook is not opened read-only.
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Public"] if "Public" in wb.sheetnames else wb.worksheets[0]
    hdr_row = next(r for r in range(1, 6) if any(str(c.value or "").strip() == "Company" for c in ws[r]))
    # Header cells carry manual line breaks ("Incident\nDate").
    hdr = [" ".join(str(c.value or "").split()) for c in ws[hdr_row]]
    col = {h: i for i, h in enumerate(hdr)}
    rows = []
    for r in ws.iter_rows(min_row=hdr_row + 1):
        vals = [c.value for c in r]
        if not any(v not in (None, "") for v in vals[1:]):
            continue
        date = vals[col["Incident Date"]]
        rep = r[col["Report"]] if "Report" in col else None
        link = rep.hyperlink.target if rep is not None and rep.hyperlink else None
        flag = lambda h: str(vals[col[h]] or "").strip() not in ("", "None", "—", "-")
        rows.append({"date": date.date().isoformat() if isinstance(date, datetime.datetime) else clean(date),
                     "company": clean(vals[col["Company"]]), "city": clean(vals[col["City"]]),
                     "state": clean(vals[col["State"]]),
                     "fatality": flag("Fatality"), "serious_injury": flag("Serious Injury"),
                     "property_damage": flag("Substantial Property Damage"),
                     "report": urllib.parse.urljoin(CSB, link) if link else None})
    return rows, url, vols


def cmd_csb_releases(a):
    rows, url, vols = arre_rows(a.refresh)
    total = len(rows)
    if a.state:
        st = STATES.get(a.state.upper(), a.state) if len(a.state) == 2 else a.state
        rows = [r for r in rows if (r["state"] or "").lower() == st.lower()]
    if a.company:
        rows = [r for r in rows if words_in(a.company, r["company"])]
    if a.city:
        rows = [r for r in rows if words_in(a.city, r["city"])]
    if a.csv:
        with open(a.csv, "w", newline="", encoding="utf8") as fh:
            w = csv.DictWriter(fh, fieldnames=list(rows[0]) if rows else ["date"])
            w.writeheader()
            w.writerows(rows)
    if not rows:
        die("no Accidental Release Events match (of %d in the workbook)" % total, 3)
    return {"count": len(rows), "total": total, "rows": rows, "workbook": url, "volumes": vols,
            "page": ARRE_PAGE, "retrieved": today(), "csv": a.csv, "note": ARRE_NOTE}


# ---------------------------------------------------------------- health

def cmd_health(a):
    out = {"checked": today(), "checks": []}

    def check(name, fn):
        try:
            out["checks"].append({"name": name, "ok": True, "detail": fn()})
        except (SourceDown, StopIteration, KeyError) as e:
            out["checks"].append({"name": name, "ok": False, "detail": str(e)})

    y = datetime.date.today().year

    def nrc_cur():
        h = head(NRC_FILE % (y % 100))
        if h is None:
            raise SourceDown("CY%02d.xlsx not found" % (y % 100))
        return "CY%d: %.1f MB, published %s" % (y, h[0] / 1e6, h[1])

    def nrc_old():
        h = head(NRC_FILE % (NRC_FIRST % 100))
        if h is None:
            raise SourceDown("CY90.xlsx not found")
        return "CY%d: %.1f MB" % (NRC_FIRST, h[0] / 1e6)

    def csb_list():
        n = len(csb_cards("current", True))
        if not n:
            raise SourceDown("no investigation cards parsed; markup may have changed")
        return "%d current investigations parsed" % n

    def arre():
        url, vols = arre_link(True)
        return "%s (+%d compiled volumes)" % (url, len(vols))

    check("NRC current-year workbook", nrc_cur)
    check("NRC 1990 workbook", nrc_old)
    check("CSB current investigations", csb_list)
    check("CSB Accidental Release Events link", arre)
    idx = nrc_index()
    out["nrc_cached"] = {k: "%s reports, published %s" % (v.get("reports"), v.get("last_modified"))
                         for k, v in sorted(idx.items())}
    try:
        import openpyxl  # noqa: F401
        out["openpyxl"] = True
    except ImportError:
        out["openpyxl"] = False
    return out


# ---------------------------------------------------------------- output

def fmt_mat(m):
    if "continuous_upper_bound" in m:
        return "%s (continuous release, upper bound %s %s %s)" % (m.get("name"), m.get("continuous_upper_bound"),
                                                                  m.get("unit") or "", m.get("rate") or "")
    amt = "%s %s" % (m["amount"], m["unit"] or "") if m.get("amount") else "unknown amount"
    s = "%s, %s" % (m.get("name") or "(unnamed)", amt.strip())
    if m.get("cas"):
        s += ", CAS %s" % m["cas"]
    if m.get("reached_water") == "YES":
        s += ", reached water" + (" (%s %s)" % (m["amount_in_water"], m.get("unit_in_water") or "")
                                  if m.get("amount_in_water") else "")
    return s


def show_nrc_row(r, long=False):
    flags = []
    if r["potential"]:
        flags.append("POTENTIAL")
    if r["update_of"]:
        flags.append("update of %s" % r["update_of"])
    if r.get("calltype") and r["calltype"] != "INC":
        flags.append(r["calltype"])
    harm = ", ".join("%s %s" % (r[k], k) for k in ("injured", "hospitalized", "fatalities", "evacuated")
                     if r.get(k) not in (None, 0, "0"))
    print("NRC %d  %s  %s / %s%s" % (r["seqnos"], r["incident_date"] or r["received"], r["type"], r["cause"],
                                    "  [%s]" % "; ".join(flags) if flags else ""))
    print("  %s | %s, %s %s%s" % (r["company"] or "(no responsible party given)",
                                  r["address"] or r["location"] or "", r["city"] or "", r["county"] or "",
                                  "  (%.2f km)" % r["distance_km"] if "distance_km" in r else ""))
    for m in r["materials"]:
        print("  - " + fmt_mat(m))
    if r["medium"] or harm:
        print("  medium: %s%s%s" % (r["medium"], " (%s)" % r["body_of_water"] if r["body_of_water"] else "",
                                   "; " + harm if harm else ""))
    d = r["description"] or ""
    if not long:
        d = " ".join(d.split())
        d = d if len(d) < 300 else d[:300] + " ..."
    print("  " + d.replace("\n", "\n  "))


def show_vintage(v):
    for x in v:
        print("  CY%d workbook, Last-Modified %s, retrieved %s: %s" % (x["year"], x["published"], x["retrieved"],
                                                                     x["url"]))


def show(cmd, r):
    if cmd == "health":
        for c in r["checks"]:
            print("%-38s %s  %s" % (c["name"], "OK  " if c["ok"] else "FAIL", c["detail"]))
        print("openpyxl available: %s" % r["openpyxl"])
        if r["nrc_cached"]:
            print("NRC years cached:")
            for k, v in r["nrc_cached"].items():
                print("  CY%s  %s" % (k, v))
        return
    if cmd == "nrc":
        print("NRC reports matching %s: %d" % (r["filters"], r["count"]))
        if r["count"]:
            print("  of which: updates or corrections to an earlier report %d; potential, no confirmed release %d"
                  % (r["updates"], r["potential"]))
            if r.get("no_coords_skipped"):
                print("  %d otherwise-matching reports had no coordinates and were left out of the radius search"
                      % r["no_coords_skipped"])
            print()
            for row in r["rows"]:
                show_nrc_row(row)
                print()
        if r.get("csv"):
            print("CSV written: %s" % r["csv"])
        print("Source: U.S. Coast Guard National Response Center FOIA data")
        show_vintage(r["vintage"])
        if r.get("note"):
            print("\n" + r["note"])
        return
    if cmd == "nrc-summary":
        print("NRC reports matching %s: %d (%d updates, %d potential)" % (r["filters"], r["count"], r["updates"],
                                                                          r["potential"]))
        t = r["totals"]
        print("Reported harm (caller's figures): %d injured, %d hospitalized, %d fatalities, %d evacuated; "
              "%d reports with injuries, %d with evacuations, %d with material reaching water"
              % (t.get("injured", 0), t.get("hospitalized", 0), t.get("fatalities", 0), t.get("evacuated", 0),
                 t.get("with_injury", 0), t.get("with_evacuation", 0), t.get("reached_water", 0)))
        for k in ("year", "type", "medium", "county", "cause", "company", "material"):
            print("\nBy %s:" % k)
            for name, n in r["by"][k]:
                print("  %5d  %s" % (n, name))
        print("\nSource: U.S. Coast Guard National Response Center FOIA data")
        show_vintage(r["vintage"])
        print("\n" + r["note"])
        return
    if cmd == "nrc-report":
        show_nrc_row(r["summary"], long=True)
        for sheet, rows in r["sheets"].items():
            if sheet == "MATERIAL_INVOLVED":
                continue
            for row in rows:
                print("\n[%s]" % sheet)
                for k, v in row.items():
                    if k == "DESCRIPTION_OF_INCIDENT":
                        continue
                    print("  %s: %s" % (k, v))
        print("\nSource: U.S. Coast Guard National Response Center FOIA data")
        show_vintage(r["vintage"])
        print("\n" + r["note"])
        return
    if cmd == "csb-investigations":
        print("CSB investigations: %d" % r["count"])
        for x in r["rows"]:
            print("\n%s  [%s]  %s" % (x["title"], x["status"], x["location"]))
            print("  %s %s" % (x["date_label"] or "", x["date"] or ""))
            print("  %s" % x["url"])
            if x["summary"]:
                print("  " + x["summary"])
        print("\nSource: %s (retrieved %s)" % (r["source"], r["retrieved"]))
        return
    if cmd == "csb-detail":
        print(r["title"])
        for k, v in r["facts"].items():
            print("  %s: %s" % (k, v))
        if r["description"]:
            print("\n" + r["description"])
        for name, links in r["documents"].items():
            print("\n%s:" % name)
            for l in links:
                print("  %s  %s" % (l["label"], l["url"]))
        if r["recommendations"]:
            print("\nRecommendations (%d):" % len(r["recommendations"]))
            for x in r["recommendations"]:
                print("  %s  to %s  [%s]" % (x["number"], x["recipient"], x["status"]))
                print("    " + (x["text"] or "").replace("\n", "\n    "))
        print("\nSource: %s (retrieved %s)" % (r["url"], r["retrieved"]))
        return
    if cmd == "csb-releases":
        print("CSB Accidental Release Events: %d of %d" % (r["count"], r["total"]))
        for x in r["rows"]:
            harm = ", ".join(k for k, f in (("fatality", x["fatality"]), ("serious injury", x["serious_injury"]),
                                            ("substantial property damage", x["property_damage"])) if f)
            print("  %s  %-28s %s, %s  [%s]" % (x["date"], x["company"], x["city"], x["state"], harm))
            if x["report"]:
                print("      %s" % x["report"])
        if r.get("csv"):
            print("CSV written: %s" % r["csv"])
        print("\nSource: CSB Accidental Release Events workbook %s, linked from %s (retrieved %s)"
              % (r["workbook"], r["page"], r["retrieved"]))
        if r["volumes"]:
            print("Compiled incident report volumes: " + ", ".join(r["volumes"]))
        print("\n" + r["note"])
        return


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--json", action="store_true", help="machine-readable output")
    ap.add_argument("--refresh", action="store_true", help="bypass the cache")
    ap.add_argument("--yes", action="store_true", help="allow an NRC download over the size guard")
    # Repeat the globals on each subcommand so they work on either side of it.
    common = argparse.ArgumentParser(add_help=False)
    common.add_argument("--json", action="store_true", default=argparse.SUPPRESS)
    common.add_argument("--refresh", action="store_true", default=argparse.SUPPRESS)
    common.add_argument("--yes", action="store_true", default=argparse.SUPPRESS)
    sub = ap.add_subparsers(dest="cmd", required=True)
    _add = sub.add_parser
    sub.add_parser = lambda *a, **k: _add(*a, parents=[common], **k)

    sub.add_parser("health", help="sources up, cached NRC years")

    def nrc_filters(s):
        s.add_argument("--state", default="LA", help="incident location state, two-letter")
        s.add_argument("--county", help="county or parish, e.g. 'St. Charles'")
        s.add_argument("--city", help="nearest city as the caller gave it")
        s.add_argument("--type", help="incident type: FIXED, PIPELINE, STORAGE TANK, VESSEL, RAILROAD, MOBILE, ...")
        s.add_argument("--company", help="responsible party words; alternatives separated by |")
        s.add_argument("--text", help="words in the description, location or address")
        s.add_argument("--material", help="material name words or exact CAS number")
        s.add_argument("--years", help="year NRC received the call, e.g. 2021-2026 or 2019,2023 (default: last 5); "
                                       "late-December incidents can sit in the next year's file")
        s.add_argument("--near", nargs=2, type=float, metavar=("LAT", "LON"))
        s.add_argument("--radius-km", type=float)
        s.add_argument("--include-drills", action="store_true")

    s = sub.add_parser("nrc", help="NRC reports by location, company, material")
    nrc_filters(s)
    s.add_argument("--csv")
    s = sub.add_parser("nrc-summary", help="counts and totals for matching NRC reports")
    nrc_filters(s)
    s = sub.add_parser("nrc-report", help="every field of one NRC report")
    s.add_argument("seqnos")
    s.add_argument("--year", type=int, help="year NRC received the call")
    s = sub.add_parser("csb-investigations", help="CSB current and completed investigations")
    s.add_argument("--state", help="two-letter code or full name")
    s.add_argument("--company", help="words in the title or summary")
    s.add_argument("--status", choices=["current", "completed"])
    s = sub.add_parser("csb-detail", help="one CSB investigation: facts, reports, recommendations")
    s.add_argument("slug", help="page slug such as dow-louisiana-operations-explosions, or its URL")
    s = sub.add_parser("csb-releases", help="Accidental Release Events reported to CSB under 40 CFR 1604")
    s.add_argument("--state", help="full name or two-letter code")
    s.add_argument("--company")
    s.add_argument("--city")
    s.add_argument("--csv")

    args = ap.parse_args()
    fn = {"health": cmd_health, "nrc": cmd_nrc, "nrc-summary": cmd_nrc_summary, "nrc-report": cmd_nrc_report,
          "csb-investigations": cmd_csb_investigations, "csb-detail": cmd_csb_detail,
          "csb-releases": cmd_csb_releases}[args.cmd]
    try:
        res = fn(args)
    except SourceDown as e:
        die("source unavailable: %s" % e, 5)
    except NeedsConfirm as e:
        die(str(e), 4)
    if args.json:
        json.dump(res, sys.stdout, indent=1, default=str)
        print()
    else:
        show(args.cmd, res)


if __name__ == "__main__":
    main()
